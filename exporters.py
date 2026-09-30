import io
import os
import re
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.lib import colors


def sanitize_text(text: str) -> str:
    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2026": "...",
        "\u00a0": " ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"[^\x09\x0A\x0D\x20-\x7E]", "", text)
    return text.strip()


def split_sections(text: str):
    lines = [line.strip() for line in sanitize_text(text).splitlines()]
    return [line for line in lines if line]


def format_txt(text: str) -> bytes:
    return (sanitize_text(text) + "\n").encode("utf-8")


def format_docx(text: str, doc_type: str, logo_path: Optional[str] = None) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    if logo_path and Path(logo_path).exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(logo_path, width=Inches(1.1))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for line in split_sections(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15

        if re.match(r"^\d+[\.\)]\s+[A-Z][A-Z0-9 /&-]+$", line):
            r = p.add_run(line)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
        elif re.match(r"^\d+[\.\)]\s+", line):
            r = p.add_run(line)
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
        else:
            r = p.add_run(line)
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("LegalEase — Draft for Review")
    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(9)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def format_pdf(text: str, doc_type: str, logo_path: Optional[str] = None) -> bytes:
    out = io.BytesIO()

    def header_footer(canvas, doc):
        canvas.saveState()
        if logo_path and Path(logo_path).exists():
            try:
                canvas.drawImage(
                    logo_path,
                    20 * mm,
                    A4[1] - 25 * mm,
                    width=16 * mm,
                    height=16 * mm,
                    preserveAspectRatio=True,
                    mask="auto",
                )
            except Exception:
                pass
        canvas.setFont("Helvetica", 8)
        canvas.drawCentredString(
            A4[0] / 2,
            12 * mm,
            f"LegalEase — Draft for Review — Page {doc.page}",
        )
        canvas.restoreState()

    doc = SimpleDocTemplate(
        out,
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=28 * mm,
        bottomMargin=20 * mm,
        title=doc_type,
        author="LegalEase",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "LegalTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        alignment=1,
        spaceAfter=14,
    )
    heading_style = ParagraphStyle(
        "LegalHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        spaceBefore=8,
        spaceAfter=5,
    )
    body_style = ParagraphStyle(
        "LegalBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        spaceAfter=7,
    )

    story = [Paragraph(doc_type.upper(), title_style)]

    for line in split_sections(text):
        if re.match(r"^\d+[\.\)]\s+[A-Z][A-Z0-9 /&-]+$", line):
            story.append(Paragraph(line, heading_style))
        else:
            escaped = (
                line.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )
            story.append(Paragraph(escaped, body_style))

    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)
    return out.getvalue()
